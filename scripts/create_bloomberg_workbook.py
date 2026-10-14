"""Create the Bloomberg Excel add-in workbook for the locked Agency MBS pools.

Bloomberg formulas cannot be evaluated outside a logged-in Terminal.  This
script therefore performs structural validation only; the user must open the
output in Terminal Excel, refresh, cross-check the pilot CUSIP with FLDS/DES,
and save the calculated workbook before Python ingestion.
"""

from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from select_securities import load_security_config  # noqa: E402

OUTPUT = ROOT / "data" / "templates" / "agency_mbs_bloomberg_pull.xlsx"
FIELDS_PATH = ROOT / "config" / "bloomberg_fields.yaml"

NAVY = "17365D"
BLUE = "0000FF"
GREEN = "008000"
RED = "FF0000"
LIGHT_BLUE = "D9EAF7"
LIGHT_GREEN = "E2F0D9"
YELLOW = "FFF2CC"
LIGHT_RED = "FCE4D6"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
THIN_GRAY = Side(style="thin", color="B7B7B7")


def _title(sheet, title: str, subtitle: str) -> None:
    sheet.merge_cells("A1:H1")
    sheet["A1"] = title
    sheet["A1"].font = Font(name="Arial", size=16, bold=True, color=WHITE)
    sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    sheet["A1"].alignment = Alignment(horizontal="left")
    sheet.merge_cells("A2:H2")
    sheet["A2"] = subtitle
    sheet["A2"].font = Font(name="Arial", size=10, italic=True, color="404040")
    sheet["A2"].alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[1].height = 24
    sheet.row_dimensions[2].height = 34


def _header(cells: Iterable) -> None:
    for cell in cells:
        cell.font = Font(name="Arial", bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GRAY)


def _input(cell, value, number_format: str | None = None) -> None:
    cell.value = value
    cell.font = Font(name="Arial", color=BLUE)
    cell.fill = PatternFill("solid", fgColor=YELLOW)
    if number_format:
        cell.number_format = number_format


def _formula(cell, formula: str, external: bool = True) -> None:
    cell.value = formula
    cell.font = Font(name="Arial", color=GREEN if not external else RED)


def _set_standard_dimensions(sheet, widths: Dict[str, float]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows():
        for cell in row:
            cell.font = cell.font.copy(name="Arial")


def _load_fields() -> Dict:
    with open(FIELDS_PATH, "r", encoding="utf-8") as stream:
        return yaml.safe_load(stream)


def _instructions(workbook: Workbook) -> None:
    sheet = workbook.active
    sheet.title = "Instructions"
    _title(
        sheet,
        "Agency MBS Bloomberg Data Pull",
        "Run this workbook only in desktop Excel with a logged-in Bloomberg add-in.",
    )
    steps = [
        ("1", "Open Inputs", "Confirm the three locked CUSIPs and choose the history dates."),
        (
            "2",
            "Validate the pilot",
            "On Bloomberg, run FLDS <GO> for 3140QVLM4 Mtge and verify every required mnemonic on Field_Check.",
        ),
        (
            "3",
            "Refresh Bloomberg",
            "Bloomberg ribbon → Refresh Workbook. Wait until no cells show Requesting Data.",
        ),
        (
            "4",
            "Cross-check",
            "Compare coupon/factor/date to DES and BVAL. Compare factor history to MTG <GO>.",
        ),
        (
            "5",
            "Review quality",
            "Data_Quality must show PASS for CUSIP, coupon, factor, and factor date.",
        ),
        (
            "6",
            "Save calculated copy",
            "Save under data/market/raw/bloomberg/. Do not commit licensed Bloomberg values.",
        ),
        (
            "7",
            "Import",
            "Run: python src/market_data.py data/market/raw/bloomberg/<calculated>.xlsx",
        ),
    ]
    sheet.append([])
    sheet.append(["Step", "Action", "Details"])
    _header(sheet[4])
    for step in steps:
        sheet.append(step)
    sheet["A13"] = "Important"
    sheet["A13"].font = Font(name="Arial", bold=True, color=RED)
    sheet["B13"] = (
        "BVAL is an evaluated price, not an executed trade. PX_LAST is diagnostic only. "
        "Frequency must be D, W, or M. Bloomberg rejects the word DAILY. "
        "BVAL_MID_PRICE is Field Not Applicable on these CL pools; use bid/ask. "
        "Weighted-average loan age is not a Bloomberg field on these pools; use the issuer file. "
        "Open this new template and refresh it. Do not reuse a previously calculated copy."
    )
    sheet["B13"].alignment = Alignment(wrap_text=True)
    sheet["B13"].fill = PatternFill("solid", fgColor=LIGHT_RED)
    sheet.merge_cells("B13:H14")
    sheet.row_dimensions[13].height = 48
    _set_standard_dimensions(sheet, {"A": 10, "B": 25, "C": 90})
    sheet.freeze_panes = "A4"


def _inputs(workbook: Workbook, securities: List[Dict], periodicity: str) -> None:
    sheet = workbook.create_sheet("Inputs")
    _title(
        sheet,
        "Inputs",
        "Blue/yellow cells are editable. Frequency must be the BDH code D, W, or M.",
    )
    headers = [
        "Role",
        "Agency",
        "Pool Number",
        "CUSIP",
        "Bloomberg Security",
        "Start Date",
        "End Date",
        "Frequency",
    ]
    for column, header in enumerate(headers, 1):
        sheet.cell(4, column, header)
    _header(sheet[4])
    for row_index, security in enumerate(securities, 5):
        values = [
            security["role"],
            security["agency"],
            security["pool_number"],
            security["cusip"],
            security["bloomberg_security"],
            security["issue_date"],
            "=TODAY()",
            periodicity,
        ]
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_index, column)
            _input(cell, value)
        sheet.cell(row_index, 6).number_format = "yyyy-mm-dd"
        sheet.cell(row_index, 7).number_format = "yyyy-mm-dd"
    validation = DataValidation(type="list", formula1='"D,W,M"', allow_blank=False)
    sheet.add_data_validation(validation)
    validation.add(f"H5:H{4 + len(securities)}")
    sheet["A10"] = "Do not change CUSIPs after validation starts."
    sheet["A10"].font = Font(name="Arial", bold=True, color=RED)
    sheet.merge_cells("A10:H10")
    sheet["A11"] = "Frequency codes: D = daily, W = weekly, M = monthly. Do not enter DAILY."
    sheet["A11"].font = Font(name="Arial", bold=True, color=RED)
    sheet.merge_cells("A11:H11")
    _set_standard_dimensions(
        sheet,
        {"A": 13, "B": 12, "C": 16, "D": 14, "E": 24, "F": 14, "G": 14, "H": 13},
    )
    sheet.freeze_panes = "A5"


def _bdh_formula(security_cell: str, mnemonic: str, start_cell: str, end_cell: str, per_cell: str) -> str:
    """BDH override string. Per must be D/W/M/Q/S/Y. Days=T is not sent."""
    return (
        f'=BDH({security_cell},"{mnemonic}",{start_cell},{end_cell},'
        f'"Per="&{per_cell},"Dir=V","Dts=S","Sort=D")'
    )


def _bloomberg_na(cell_ref: str) -> str:
    return (
        f'OR(ISERROR({cell_ref}),'
        f'ISNUMBER(SEARCH("Field Not Applicable",{cell_ref}&"")),'
        f'ISNUMBER(SEARCH("Invalid Field",{cell_ref}&"")))'
    )


def _error_status(cell_ref: str, required: bool) -> str:
    if required:
        return (
            f'=IF({_bloomberg_na(cell_ref)},"ERROR",'
            f'IF({cell_ref}="","MISSING","RETURNED"))'
        )
    return (
        f'=IF({_bloomberg_na(cell_ref)},"NOT APPLICABLE",'
        f'IF(LEFT({cell_ref}&"",1)="#","ERROR",'
        f'IF({cell_ref}="","MISSING","RETURNED")))'
    )


def _history_status(date_ref: str, value_ref: str, required: bool) -> str:
    if required:
        return (
            f'=IF({_bloomberg_na(date_ref)},"ERROR",'
            f'IF(ISNUMBER({value_ref}),"RETURNED","MISSING"))'
        )
    return (
        f'=IF({_bloomberg_na(date_ref)},"NOT APPLICABLE",'
        f'IF(LEFT({date_ref}&"",1)="#","ERROR",'
        f'IF(ISNUMBER({value_ref}),"RETURNED","MISSING")))'
    )


def _field_check(workbook: Workbook, fields: Dict) -> None:
    sheet = workbook.create_sheet("Field_Check")
    _title(
        sheet,
        "Pilot field validation",
        "Required fields must return a value for the primary CUSIP. Confirm each mnemonic with FLDS <GO>.",
    )
    headers = ["Category", "Mnemonic", "Label", "Required", "Pilot Formula", "Status", "Terminal Check"]
    for column, header in enumerate(headers, 1):
        sheet.cell(4, column, header)
    _header(sheet[4])
    row = 5
    history_index = 0
    for category in ("static_fields", "historical_fields", "bulk_fields"):
        for field in fields[category]:
            sheet.cell(row, 1, category)
            sheet.cell(row, 2, field["mnemonic"])
            sheet.cell(row, 3, field["label"])
            sheet.cell(row, 4, "YES" if field["required"] else "NO")
            if field["function"] == "BDP":
                _formula(sheet.cell(row, 5), f'=BDP(Inputs!$E$5,"{field["mnemonic"]}")')
                _formula(
                    sheet.cell(row, 6),
                    _error_status(f"E{row}", field["required"]),
                    external=False,
                )
            elif field["function"] == "BDH":
                # Do not place BDH here. A live BDH spills into the Status column.
                column = 1 + history_index * 3
                date_ref = f"Price_History!{get_column_letter(column)}6"
                value_ref = f"Price_History!{get_column_letter(column + 1)}6"
                _formula(
                    sheet.cell(row, 5),
                    f'=IF(ISNUMBER({value_ref}),{value_ref},{date_ref})',
                    external=False,
                )
                _formula(
                    sheet.cell(row, 6),
                    _history_status(date_ref, value_ref, field["required"]),
                    external=False,
                )
                history_index += 1
            else:
                date_ref = "Factor_History!A5"
                value_ref = "Factor_History!B5"
                _formula(
                    sheet.cell(row, 5),
                    f'=IF(ISNUMBER({value_ref}),{value_ref},{date_ref})',
                    external=False,
                )
                _formula(
                    sheet.cell(row, 6),
                    _history_status(date_ref, value_ref, field["required"]),
                    external=False,
                )
            sheet.cell(row, 7, "Confirm in FLDS")
            sheet.cell(row, 7).fill = PatternFill("solid", fgColor=YELLOW)
            row += 1
    sheet.conditional_formatting.add(
        f"F5:F{row}",
        FormulaRule(formula=['F5="ERROR"'], fill=PatternFill("solid", fgColor=LIGHT_RED)),
    )
    _set_standard_dimensions(
        sheet, {"A": 18, "B": 24, "C": 34, "D": 11, "E": 72, "F": 13, "G": 20}
    )
    sheet.freeze_panes = "A5"


def _static_reference(workbook: Workbook, securities: List[Dict], fields: Dict) -> None:
    sheet = workbook.create_sheet("Static_Reference")
    _title(sheet, "Static security reference", "Bloomberg BDP values; issuer disclosure remains authoritative.")
    static = fields["static_fields"]
    sheet.cell(4, 1, "Role")
    sheet.cell(4, 2, "Bloomberg Security")
    for column, field in enumerate(static, 3):
        sheet.cell(4, column, field["mnemonic"])
        sheet.cell(4, column).comment = Comment(field["label"], "MBS Analytics")
    _header(sheet[4])
    for index, security in enumerate(securities, 5):
        sheet.cell(index, 1, security["role"])
        sheet.cell(index, 2, security["bloomberg_security"])
        for column, field in enumerate(static, 3):
            _formula(
                sheet.cell(index, column),
                f'=BDP($B{index},"{field["mnemonic"]}")',
            )
    _set_standard_dimensions(sheet, {"A": 13, "B": 24})
    for column in range(3, 3 + len(static)):
        sheet.column_dimensions[get_column_letter(column)].width = 20
    sheet.freeze_panes = "C5"


def _price_history(workbook: Workbook, securities: List[Dict], fields: Dict, block_height: int) -> None:
    sheet = workbook.create_sheet("Price_History")
    _title(
        sheet,
        "Historical market observations",
        "Each block spills two columns: date and value. BVAL is evaluated; bid/ask are preferred where populated.",
    )
    historical = fields["historical_fields"]
    if block_height < 800:
        raise ValueError("price_block_height is too small for daily history since 2024.")
    for security_index, security in enumerate(securities):
        start_row = 4 + security_index * block_height
        sheet.cell(start_row, 1, security["role"].upper())
        sheet.cell(start_row, 2, security["cusip"])
        sheet.cell(start_row, 3, security["bloomberg_security"])
        for cell in sheet[start_row][:3]:
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(name="Arial", bold=True, color=WHITE)
        input_row = 5 + security_index
        for field_index, field in enumerate(historical):
            column = 1 + field_index * 3
            sheet.cell(start_row + 1, column, field["mnemonic"])
            sheet.cell(start_row + 1, column + 1, field["quality_tier"])
            _header(sheet[start_row + 1][column - 1 : column + 1])
            for header_cell in sheet[start_row + 1][column - 1 : column + 1]:
                header_cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=False
                )
            anchor = sheet.cell(start_row + 2, column)
            _formula(
                anchor,
                _bdh_formula(
                    f"Inputs!$E${input_row}",
                    field["mnemonic"],
                    f"Inputs!$F${input_row}",
                    f"Inputs!$G${input_row}",
                    f"Inputs!$H${input_row}",
                ),
            )
            anchor.alignment = Alignment(wrap_text=False)
    for column in range(1, 1 + len(historical) * 3):
        width = 3 if column % 3 == 0 else 22
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A6"


def _factor_history(workbook: Workbook, securities: List[Dict]) -> None:
    sheet = workbook.create_sheet("Factor_History")
    _title(sheet, "Historical pool factors", "Bloomberg BDS MTG_HIST_FACT; reconcile every date to issuer factors.")
    for index, security in enumerate(securities):
        column = 1 + index * 4
        input_row = 5 + index
        sheet.cell(4, column, security["role"].upper())
        sheet.cell(4, column + 1, security["cusip"])
        _header(sheet[4][column - 1 : column + 1])
        _formula(
            sheet.cell(5, column),
            f'=BDS(Inputs!$E${input_row},"MTG_HIST_FACT")',
        )
        sheet.column_dimensions[get_column_letter(column)].width = 20
        sheet.column_dimensions[get_column_letter(column + 1)].width = 18
    sheet.freeze_panes = "A5"


def _data_quality(workbook: Workbook, securities: List[Dict]) -> None:
    sheet = workbook.create_sheet("Data_Quality")
    _title(
        sheet,
        "Workbook quality controls",
        "These checks are necessary but not sufficient; Python performs issuer reconciliation after import.",
    )
    headers = ["Role", "CUSIP", "Check", "Expected", "Observed", "Status"]
    for column, header in enumerate(headers, 1):
        sheet.cell(4, column, header)
    _header(sheet[4])
    row = 5
    checks = [
        ("CUSIP", "configured CUSIP", "ID_CUSIP"),
        ("Coupon", "configured coupon", "CPN"),
        ("Factor", "public discovery factor", "MTG_FACTOR"),
        ("Factor date", "latest disclosure month", "MTG_FACTOR_DT"),
    ]
    static_headers = {
        cell.value: cell.column for cell in workbook["Static_Reference"][4] if cell.value
    }
    for security_index, security in enumerate(securities, 5):
        for label, expected_label, mnemonic in checks:
            sheet.cell(row, 1, security["role"])
            sheet.cell(row, 2, security["cusip"])
            sheet.cell(row, 3, label)
            if mnemonic == "ID_CUSIP":
                expected = security["cusip"]
            elif mnemonic == "CPN":
                # Terminal CPN is percent (5.5), not the config decimal 0.055.
                expected = round(float(security["coupon"]) * 100, 6)
            elif mnemonic == "MTG_FACTOR":
                expected = security["factor"]
            else:
                expected = datetime.fromisoformat(security["discovery_as_of"])
            sheet.cell(row, 4, expected)
            if mnemonic == "CPN":
                sheet.cell(row, 4).number_format = "0.000"
            elif mnemonic == "MTG_FACTOR":
                sheet.cell(row, 4).number_format = "0.00000000"
            elif mnemonic == "MTG_FACTOR_DT":
                sheet.cell(row, 4).number_format = "yyyy-mm-dd"
            if mnemonic == "MTG_FACTOR_DT":
                factor_column = get_column_letter(1 + (security_index - 5) * 4)
                observed = f"=Factor_History!{factor_column}5"
            else:
                source_col = get_column_letter(static_headers[mnemonic])
                observed = f"=Static_Reference!{source_col}{security_index}"
            _formula(sheet.cell(row, 5), observed, external=False)
            if mnemonic == "MTG_FACTOR_DT":
                sheet.cell(row, 5).number_format = "yyyy-mm-dd"
            if mnemonic == "ID_CUSIP":
                status_formula = (
                    f'=IF(OR(ISERROR(E{row}),LEFT(E{row}&"",1)="#"),"ERROR",'
                    f'IF(E{row}=D{row},"PASS","FAIL"))'
                )
            elif mnemonic == "CPN":
                status_formula = (
                    f'=IF(OR(ISERROR(E{row}),LEFT(E{row}&"",1)="#"),"ERROR",'
                    f'IF(OR(ABS(E{row}-D{row})<=0.001,ABS(E{row}*100-D{row})<=0.001),"PASS","CHECK"))'
                )
            elif mnemonic == "MTG_FACTOR":
                status_formula = (
                    f'=IF(OR(ISERROR(E{row}),LEFT(E{row}&"",1)="#"),"ERROR",'
                    f'IF(ABS(E{row}-D{row})<=0.01,"PASS","CHECK"))'
                )
            else:
                # Bloomberg MTG_HIST_FACT dates these CL pools on the 25th,
                # not month-end. Pass when year and month match.
                status_formula = (
                    f'=IF(OR(ISERROR(E{row}),LEFT(E{row}&"",1)="#"),"ERROR",'
                    f'IF(E{row}="","MISSING",'
                    f'IF(AND(YEAR(E{row})=YEAR(D{row}),MONTH(E{row})=MONTH(D{row})),'
                    f'"PASS","CHECK")))'
                )
            _formula(sheet.cell(row, 6), status_formula, external=False)
            sheet.cell(row, 4).comment = Comment(expected_label, "MBS Analytics")
            row += 1
    sheet.conditional_formatting.add(
        f"F5:F{row}",
        FormulaRule(formula=['F5="FAIL"'], fill=PatternFill("solid", fgColor=LIGHT_RED)),
    )
    sheet.conditional_formatting.add(
        f"F5:F{row}",
        FormulaRule(formula=['F5="PASS"'], fill=PatternFill("solid", fgColor=LIGHT_GREEN)),
    )
    _set_standard_dimensions(
        sheet, {"A": 13, "B": 14, "C": 22, "D": 24, "E": 24, "F": 14}
    )
    sheet.freeze_panes = "A5"


def _export(workbook: Workbook, securities: List[Dict]) -> None:
    sheet = workbook.create_sheet("Export")
    _title(
        sheet,
        "Python import contract",
        "Do not edit. Python imports static values here and scans Price_History / Factor_History blocks.",
    )
    headers = [
        "role",
        "agency",
        "pool_number",
        "cusip",
        "bloomberg_security",
        "start_date",
        "end_date",
    ]
    for column, header in enumerate(headers, 1):
        sheet.cell(4, column, header)
    _header(sheet[4])
    for row, _security in enumerate(securities, 5):
        input_row = row
        for column, source_column in enumerate("ABCDEFG", 1):
            _formula(
                sheet.cell(row, column),
                f"=Inputs!${source_column}${input_row}",
                external=False,
            )
    for row in range(5, 5 + len(securities)):
        sheet.cell(row, 6).number_format = "yyyy-mm-dd"
        sheet.cell(row, 7).number_format = "yyyy-mm-dd"
    _set_standard_dimensions(
        sheet, {"A": 13, "B": 12, "C": 16, "D": 14, "E": 24, "F": 14, "G": 14}
    )
    sheet.freeze_panes = "A5"


def validate_workbook(path: Path) -> None:
    """Static validation for workbook structure and Bloomberg formula syntax."""
    workbook = load_workbook(path, data_only=False)
    expected = {
        "Instructions",
        "Inputs",
        "Field_Check",
        "Static_Reference",
        "Price_History",
        "Factor_History",
        "Data_Quality",
        "Export",
    }
    if set(workbook.sheetnames) != expected:
        raise ValueError(f"Unexpected sheets: {workbook.sheetnames}")
    formulas = [
        cell.value
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    ]
    if not formulas:
        raise ValueError("Workbook contains no formulas.")
    external = [formula for formula in formulas if formula.startswith(("=BDP(", "=BDH(", "=BDS("))]
    if len(external) < 30:
        raise ValueError(f"Expected at least 30 Bloomberg formulas, found {len(external)}.")
    bad = [formula for formula in formulas if "#REF!" in formula or "#NAME?" in formula]
    if bad:
        raise ValueError(f"Invalid formula references: {bad[:3]}")
    blob = "\n".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    )
    if "MTG_WALA" in blob or "Per=DAILY" in blob or "Days=T" in blob:
        raise ValueError("Workbook still contains a mnemonic or override the terminal rejected.")
    bdh = [formula for formula in formulas if formula.startswith("=BDH(")]
    if not bdh or any("Per=" not in formula or "Days=" in formula for formula in bdh):
        raise ValueError("Every BDH formula must pass Per= and must not pass Days=.")
    if any(formula.startswith("=BDH(") or formula.startswith("=BDS(") for formula in (
        cell.value for cell in workbook["Field_Check"]["E"] if isinstance(cell.value, str)
    )):
        raise ValueError("Field_Check must not contain spilling BDH or BDS formulas.")


def main() -> None:
    config = load_security_config()
    fields = _load_fields()
    securities = config["securities"]
    periodicity = str(fields["layout"]["periodicity"])
    block_height = int(fields["layout"]["price_block_height"])
    if periodicity not in {"D", "W", "M", "Q", "S", "Y"}:
        raise ValueError(f"Bloomberg Per enum does not include {periodicity}.")
    workbook = Workbook()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    _instructions(workbook)
    _inputs(workbook, securities, periodicity)
    _field_check(workbook, fields)
    _static_reference(workbook, securities, fields)
    _price_history(workbook, securities, fields, block_height)
    _factor_history(workbook, securities)
    _data_quality(workbook, securities)
    _export(workbook, securities)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    validate_workbook(OUTPUT)
    print(f"Wrote {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
