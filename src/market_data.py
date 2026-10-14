"""Import a calculated Bloomberg Terminal workbook into typed parquet files."""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import List, Tuple

import pandas as pd
import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
FIELDS_PATH = ROOT / "config" / "bloomberg_fields.yaml"


def _field_config() -> dict:
    with open(FIELDS_PATH, "r", encoding="utf-8") as stream:
        return yaml.safe_load(stream)


_CONFIG = _field_config()
BLOCK_HEIGHT = int(_CONFIG["layout"]["price_block_height"])

PRICE_FIELD_MAP = {
    "BVAL_MID_PRICE": ("bval", "evaluated"),
    "PX_BID": ("bid", "quote"),
    "PX_ASK": ("ask", "quote"),
    "PX_LAST": ("last", "diagnostic"),
    "OAS_SPREAD_MID": ("bloomberg_oas", "external_model"),
    "Z_SPRD_MID": ("bloomberg_zspread", "external_model"),
}


def _parse_excel_date(value) -> pd.Timestamp | None:
    """Accept datetime cells, ISO text, or Excel serial days from a Terminal save."""
    if value in (None, ""):
        return None
    if isinstance(value, str):
        text = value.strip()
        if not text or text.startswith("#"):
            return None
        parsed = pd.to_datetime(text, errors="coerce")
        return None if pd.isna(parsed) else pd.Timestamp(parsed).normalize()
    if isinstance(value, pd.Timestamp):
        return value.normalize()
    if hasattr(value, "year") and hasattr(value, "month") and hasattr(value, "day"):
        return pd.Timestamp(value).normalize()
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and 20_000 <= float(value) <= 80_000:
        return pd.to_datetime(float(value), unit="D", origin="1899-12-30").normalize()
    return None


def _inputs(workbook) -> pd.DataFrame:
    sheet = workbook["Inputs"]
    rows = []
    for row in range(5, 8):
        values = [sheet.cell(row, column).value for column in range(1, 9)]
        if not values[3]:
            continue
        rows.append(
            {
                "role": values[0],
                "agency": values[1],
                "pool_number": values[2],
                "cusip": str(values[3]).upper(),
                "bloomberg_security": values[4],
                "start_date": _parse_excel_date(values[5]),
                "end_date": _parse_excel_date(values[6]),
                "frequency": values[7],
            }
        )
    if len(rows) != 3:
        raise ValueError("Calculated workbook must contain exactly three configured securities.")
    return pd.DataFrame(rows)


def _historical_fields() -> List[str]:
    return [field["mnemonic"] for field in _CONFIG["historical_fields"]]


def _parse_price_history(workbook, inputs: pd.DataFrame, source_path: Path) -> pd.DataFrame:
    sheet = workbook["Price_History"]
    records = []
    fields = _historical_fields()
    for security_index, security in inputs.reset_index(drop=True).iterrows():
        start_row = 4 + security_index * BLOCK_HEIGHT
        for field_index, field in enumerate(fields):
            column = 1 + field_index * 3
            source, quality_tier = PRICE_FIELD_MAP[field]
            for row in range(start_row + 2, start_row + BLOCK_HEIGHT):
                date_value = _parse_excel_date(sheet.cell(row, column).value)
                numeric_value = sheet.cell(row, column + 1).value
                if date_value is None or numeric_value in (None, ""):
                    continue
                try:
                    value = float(numeric_value)
                except (TypeError, ValueError):
                    continue
                records.append(
                    {
                        "role": security["role"],
                        "agency": security["agency"],
                        "pool_number": security["pool_number"],
                        "cusip": security["cusip"],
                        "date": date_value,
                        "field": field,
                        "source": source,
                        "quality_tier": quality_tier,
                        "value": value,
                        "workbook": source_path.name,
                    }
                )
    return pd.DataFrame(records)


def _parse_factor_history(workbook, inputs: pd.DataFrame, source_path: Path) -> pd.DataFrame:
    sheet = workbook["Factor_History"]
    records = []
    for security_index, security in inputs.reset_index(drop=True).iterrows():
        column = 1 + security_index * 4
        for row in range(5, sheet.max_row + 1):
            date_value = _parse_excel_date(sheet.cell(row, column).value)
            factor_value = sheet.cell(row, column + 1).value
            if date_value is None or factor_value in (None, ""):
                continue
            try:
                factor = float(factor_value)
            except (TypeError, ValueError):
                continue
            records.append(
                {
                    "role": security["role"],
                    "agency": security["agency"],
                    "pool_number": security["pool_number"],
                    "cusip": security["cusip"],
                    "date": date_value,
                    "factor": factor,
                    "source": "bloomberg_mtg_hist_fact",
                    "workbook": source_path.name,
                }
            )
    return pd.DataFrame(records)


def _parse_static(workbook, inputs: pd.DataFrame, source_path: Path) -> pd.DataFrame:
    sheet = workbook["Static_Reference"]
    headers = [sheet.cell(4, column).value for column in range(1, sheet.max_column + 1)]
    rows = []
    for row in range(5, 8):
        values = [sheet.cell(row, column).value for column in range(1, sheet.max_column + 1)]
        record = dict(zip(headers, values))
        if not record.get("Bloomberg Security"):
            continue
        record["cusip_configured"] = inputs.iloc[row - 5]["cusip"]
        record["workbook"] = source_path.name
        rows.append(record)
    return pd.DataFrame(rows)


def load_bloomberg_workbook(
    path: Path,
    require_market_data: bool = True,
) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Return static reference, long market history and factor history."""
    workbook = load_workbook(path, data_only=True, read_only=False)
    inputs = _inputs(workbook)
    market = _parse_price_history(workbook, inputs, path)
    factors = _parse_factor_history(workbook, inputs, path)
    static = _parse_static(workbook, inputs, path)
    if require_market_data and market.empty:
        raise ValueError(
            "No cached Bloomberg history found. Open this workbook in Terminal Excel, "
            "refresh Bloomberg formulas, wait for completion, save, and retry."
        )
    return static, market, factors


def build_price_observations(market_long: pd.DataFrame) -> pd.DataFrame:
    """Create source-typed clean-price observations from long Bloomberg fields."""
    if market_long.empty:
        return pd.DataFrame(
            columns=["role", "agency", "pool_number", "cusip", "date", "source", "price"]
        )
    key = ["role", "agency", "pool_number", "cusip", "date"]
    wide = market_long.pivot_table(index=key, columns="source", values="value", aggfunc="last")
    rows = []
    for _, record in wide.reset_index().iterrows():
        base = {column: record[column] for column in key}
        external = {
            "external_oas_bps": record.get("bloomberg_oas"),
            "external_zspread_bps": record.get("bloomberg_zspread"),
        }
        if pd.notna(record.get("bval")):
            rows.append(
                {**base, **external, "source": "bval", "price": float(record["bval"])}
            )
        if pd.notna(record.get("bid")) and pd.notna(record.get("ask")):
            bid, ask = float(record["bid"]), float(record["ask"])
            rows.append(
                {
                    **base,
                    **external,
                    "source": "bid_ask",
                    "price": 0.5 * (bid + ask),
                    "bid": bid,
                    "ask": ask,
                }
            )
        if pd.notna(record.get("last")):
            rows.append(
                {**base, **external, "source": "last", "price": float(record["last"])}
            )
    result = pd.DataFrame(rows)
    if result.empty:
        return result
    invalid = result.loc[~result["price"].between(50.0, 150.0)]
    if not invalid.empty:
        raise ValueError(
            "Bloomberg price observations outside 50–150; verify clean-price units: "
            f"{invalid[['cusip', 'date', 'source', 'price']].head().to_dict('records')}"
        )
    if "bid" in result.columns:
        crossed = result.loc[
            result["bid"].notna() & result["ask"].notna() & (result["bid"] > result["ask"])
        ]
        if not crossed.empty:
            raise ValueError(f"Crossed Bloomberg markets found: {len(crossed)}")
    return result.sort_values(["cusip", "date", "source"]).reset_index(drop=True)


def validate_factors(factors: pd.DataFrame, tolerance: float = 1e-8) -> None:
    """Reject invalid or materially increasing historical factors."""
    if factors.empty:
        raise ValueError("No Bloomberg factor history found.")
    if not factors["factor"].between(0, 1, inclusive="right").all():
        raise ValueError("Pool factors must be in (0, 1].")
    ordered = factors.sort_values(["cusip", "date"])
    changes = ordered.groupby("cusip")["factor"].diff()
    increases = changes[changes > tolerance]
    if not increases.empty:
        raise ValueError(f"Pool factor increases require correction review: {len(increases)}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument(
        "--output",
        type=Path,
        default=ROOT / "data" / "market" / "processed",
        help="Directory for parquet outputs.",
    )
    args = parser.parse_args()

    static, market_long, factors = load_bloomberg_workbook(args.workbook)
    prices = build_price_observations(market_long)
    validate_factors(factors)
    args.output.mkdir(parents=True, exist_ok=True)
    static.to_parquet(args.output / "bloomberg_static.parquet", index=False)
    market_long.to_parquet(args.output / "bloomberg_history_long.parquet", index=False)
    prices.to_parquet(args.output / "market_observations.parquet", index=False)
    factors.to_parquet(args.output / "bloomberg_factor_history.parquet", index=False)
    print(
        f"Wrote {len(prices):,} price observations and {len(factors):,} factors "
        f"to {args.output}"
    )


if __name__ == "__main__":
    main()
