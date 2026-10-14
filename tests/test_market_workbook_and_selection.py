from datetime import datetime
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from market_data import (
    BLOCK_HEIGHT,
    build_price_observations,
    load_bloomberg_workbook,
    validate_factors,
    _parse_excel_date,
)
from select_securities import (
    DEFAULT_CONFIG,
    cusip_is_valid,
    load_security_config,
    score_securities,
)

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "data" / "templates" / "agency_mbs_bloomberg_pull.xlsx"


def test_locked_security_config_is_valid():
    config = load_security_config(DEFAULT_CONFIG)
    assert [row["role"] for row in config["securities"]] == [
        "primary",
        "holdout",
        "holdout",
    ]
    assert all(cusip_is_valid(row["cusip"]) for row in config["securities"])
    report = score_securities(config["securities"])
    assert report["disclosure_score"].eq(100).all()
    assert not report["go_no_go"].any()


def test_market_gate_passes_on_dense_quotes_without_bval():
    config = load_security_config(DEFAULT_CONFIG)
    dates = pd.date_range("2025-01-01", periods=252, freq="B")
    rows = [
        {
            "cusip": security["cusip"],
            "date": date,
            "source": "bid_ask",
            "price": 100.0,
        }
        for security in config["securities"]
        for date in dates
    ]
    report = score_securities(config["securities"], pd.DataFrame(rows))
    assert report["go_no_go"].all()
    assert report["dense_bval"].eq(False).all()
    assert report["market_status"].eq("pass_quotes_no_bval").all()


def test_market_gate_requires_dense_bval_and_observed_market():
    config = load_security_config(DEFAULT_CONFIG)
    dates = pd.date_range("2025-01-01", periods=252, freq="B")
    rows = []
    for security in config["securities"]:
        rows.extend(
            {
                "cusip": security["cusip"],
                "date": date,
                "source": "bval",
                "price": 100.0,
            }
            for date in dates
        )
        rows.extend(
            {
                "cusip": security["cusip"],
                "date": date,
                "source": "bid_ask",
                "price": 100.0,
            }
            for date in dates[:20]
        )
    report = score_securities(config["securities"], pd.DataFrame(rows))
    assert report["go_no_go"].all()


def test_excel_serial_dates_from_terminal_saves():
    assert _parse_excel_date(46304) == pd.Timestamp("2026-10-09")
    assert _parse_excel_date("2024-09-01") == pd.Timestamp("2024-09-01")
    assert _parse_excel_date("#N/A Field Not Applicable") is None


def test_template_requires_terminal_refresh():
    with pytest.raises(ValueError, match="No cached Bloomberg history"):
        load_bloomberg_workbook(TEMPLATE)


def test_import_mock_calculated_workbook(tmp_path):
    calculated = tmp_path / "calculated.xlsx"
    workbook = load_workbook(TEMPLATE)
    for row in range(5, 8):
        workbook["Inputs"].cell(row, 7, datetime(2026, 10, 10))
    # BVAL anchor is two rows below each CUSIP block header.
    for index, price in enumerate((99.5, 98.5, 101.25)):
        row = 6 + index * BLOCK_HEIGHT
        workbook["Price_History"].cell(row, 1, datetime(2026, 10, 9))
        workbook["Price_History"].cell(row, 2, price)
    # One historical factor row per security.
    for column, factor in ((1, 0.87), (5, 0.90), (9, 0.62)):
        workbook["Factor_History"].cell(5, column, datetime(2026, 10, 1))
        workbook["Factor_History"].cell(5, column + 1, factor)
    workbook.save(calculated)

    _static, market, factors = load_bloomberg_workbook(calculated)
    assert len(market) == 3
    assert len(factors) == 3
    prices = build_price_observations(market)
    assert set(prices["source"]) == {"bval"}
    validate_factors(factors)


def test_workbook_matches_terminal_field_results():
    workbook = load_workbook(TEMPLATE)
    assert workbook["Inputs"]["H5"].value == "D"
    assert [workbook["Inputs"].cell(row, 8).value for row in range(5, 8)] == ["D", "D", "D"]
    price_formula = workbook["Price_History"]["A6"].value
    assert price_formula == (
        '=BDH(Inputs!$E$5,"BVAL_MID_PRICE",Inputs!$F$5,Inputs!$G$5,'
        '"Per="&Inputs!$H$5,"Dir=V","Dts=S","Sort=D")'
    )
    assert "Days=" not in price_formula
    assert workbook["Data_Quality"]["D6"].value == 5.5
    assert workbook["Data_Quality"]["D10"].value == 5.0
    assert workbook["Data_Quality"]["D14"].value == 6.0
    assert workbook["Data_Quality"]["E8"].value == "=Factor_History!A5"
    assert workbook["Data_Quality"]["E12"].value == "=Factor_History!E5"
    assert workbook["Data_Quality"]["E16"].value == "=Factor_History!I5"
    assert "YEAR(E8)=YEAR(D8)" in workbook["Data_Quality"]["F8"].value
    bval_row = next(
        row
        for row in range(5, 25)
        if workbook["Field_Check"].cell(row, 2).value == "BVAL_MID_PRICE"
    )
    assert workbook["Field_Check"].cell(bval_row, 4).value == "NO"
    assert "Field Not Applicable" in workbook["Field_Check"].cell(bval_row, 6).value
    assert workbook["Export"]["F5"].number_format == "yyyy-mm-dd"
    assert workbook["Export"]["G5"].number_format == "yyyy-mm-dd"
    headers = [cell.value for cell in workbook["Static_Reference"][4]]
    assert "MTG_WALA" not in headers
    assert "MTG_WAM" in headers
    field_checks = [
        cell.value
        for cell in workbook["Field_Check"]["E"]
        if isinstance(cell.value, str)
    ]
    assert not any(value.startswith(("=BDH(", "=BDS(")) for value in field_checks)
    second = workbook["Price_History"].cell(6 + BLOCK_HEIGHT, 1).value
    assert "Inputs!$E$6" in second
